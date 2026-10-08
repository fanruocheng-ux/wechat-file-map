# -*- coding: utf-8 -*-
"""
微信文件地图 —— 本地微信文件夹索引 / 搜索 / 可视化工具
版权所有 © 2026 FRC翡冷翠 · 保留所有权利
只读工具：绝不移动、改名、修改、删除微信原文件；所有衍生数据只写本项目 data/ 目录。
启动: python server.py  →  http://127.0.0.1:8012
"""
import os
import re
import io
import sys
import time
import string
import hashlib
import sqlite3
import threading
import traceback
import mimetypes
from datetime import date, timedelta

from flask import Flask, request, jsonify, send_file, abort

# ---------------------------------------------------------------- 常量
FROZEN = getattr(sys, "frozen", False)                      # PyInstaller 打包后为 True
BASE_DIR = os.path.dirname(sys.executable) if FROZEN else os.path.dirname(os.path.abspath(__file__))
MODULE_DIR = os.path.dirname(os.path.abspath(__file__))     # 打包后 = 解包目录(_MEIPASS)，静态资源随之
DATA_DIR = os.path.join(BASE_DIR, "data")                   # 索引/缓存/导出 → 永远放 exe 旁边
DB_PATH = os.path.join(DATA_DIR, "index.db")
THUMB_DIR = os.path.join(DATA_DIR, "thumb_cache")
EXPORT_DIR = os.path.join(DATA_DIR, "exports")
STATIC_DIR = os.path.join(MODULE_DIR, "static")
PORT = 8012

# ---- 版权信息（署名/年份在这里改一处即可） ----
VERSION = "1.1.0"
AUTHOR = "FRC翡冷翠"
APP_NAME = "微信文件地图"
COPYRIGHT = f"© 2026 {AUTHOR} · 保留所有权利"

HASH_LIMIT = 30 * 1024 * 1024       # ≤30MB 才计算 MD5（46万文件实测：100MB 太慢；超大视频去重价值低）
CONTENT_LIMIT = 200 * 1024          # 正文最多入库字节数
EXTRACT_PAGES = 8                   # PDF 最多提取前 8 页
SNIPPET_LEN = 130                   # 内容命中摘要长度

TEXT_EXTS = {"txt", "md", "markdown", "log", "csv", "json", "xml", "html", "htm",
             "ini", "cfg", "conf", "srt", "lrc", "py", "js", "css", "sql", "bat", "yaml", "yml"}
DOC_EXTS = {"pdf", "doc", "docx", "ppt", "pptx", "txt", "md", "rtf"}
SHEET_EXTS = {"xls", "xlsx", "xlsm", "csv"}

CATEGORY_MAP = [
    ("图片", {"jpg", "jpeg", "png", "gif", "bmp", "webp", "heic", "heif", "tiff", "tif", "svg", "ico"}),
    ("视频", {"mp4", "avi", "mov", "mkv", "flv", "wmv", "ts", "3gp", "m4v", "webm", "mpg", "mpeg"}),
    ("音频", {"mp3", "wav", "amr", "aac", "m4a", "flac", "ogg", "wma", "ape"}),
    ("表格", {"xls", "xlsx", "xlsm", "csv"}),
    ("文档", {"pdf", "doc", "docx", "ppt", "pptx", "txt", "md", "rtf"}),
    ("压缩包", {"zip", "rar", "7z", "tar", "gz", "bz2", "xz", "iso"}),
]


def category_of(ext):
    for cat, exts in CATEGORY_MAP:
        if ext in exts:
            return cat
    return "其他"


# ---------------------------------------------------------------- 数据库
SCHEMA = """
CREATE TABLE IF NOT EXISTS roots(
  id INTEGER PRIMARY KEY AUTOINCREMENT,
  path TEXT UNIQUE NOT NULL,
  added_at REAL NOT NULL
);
CREATE TABLE IF NOT EXISTS files(
  id INTEGER PRIMARY KEY AUTOINCREMENT,
  root_id INTEGER NOT NULL,
  path TEXT UNIQUE NOT NULL,
  name TEXT NOT NULL,
  ext TEXT NOT NULL,
  category TEXT NOT NULL,
  size INTEGER NOT NULL,
  ctime REAL NOT NULL,
  mtime REAL NOT NULL,
  folder TEXT NOT NULL DEFAULT '',
  content TEXT,
  content_status TEXT NOT NULL DEFAULT 'pending',
  hash TEXT,
  hash_status TEXT NOT NULL DEFAULT 'pending',
  scan_ts REAL NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_files_mtime  ON files(mtime);
CREATE INDEX IF NOT EXISTS idx_files_ctime  ON files(ctime);
CREATE INDEX IF NOT EXISTS idx_files_cat    ON files(category);
CREATE INDEX IF NOT EXISTS idx_files_name   ON files(name);
CREATE INDEX IF NOT EXISTS idx_files_hash   ON files(hash);
CREATE INDEX IF NOT EXISTS idx_files_cstat  ON files(content_status);
CREATE INDEX IF NOT EXISTS idx_files_hstat  ON files(hash_status);
CREATE TABLE IF NOT EXISTS errors(
  id INTEGER PRIMARY KEY AUTOINCREMENT,
  path TEXT, reason TEXT, phase TEXT, ts REAL
);
CREATE TABLE IF NOT EXISTS meta(key TEXT PRIMARY KEY, value TEXT);
"""

# FTS5 trigram 全文索引：子串搜索从全表 LIKE（46万行 17s）降到毫秒级。
# 外部内容表 + 触发器自动同步；仅在 name/path 变化时触发，提取正文/哈希的 UPDATE 不触发。
FTS_DDL = """
CREATE VIRTUAL TABLE IF NOT EXISTS fts_name USING fts5(
  name, path, content='files', content_rowid='id', tokenize='trigram');
CREATE TRIGGER IF NOT EXISTS files_fts_ai AFTER INSERT ON files BEGIN
  INSERT INTO fts_name(rowid, name, path) VALUES (new.id, new.name, new.path);
END;
CREATE TRIGGER IF NOT EXISTS files_fts_au AFTER UPDATE OF name, path ON files BEGIN
  INSERT INTO fts_name(fts_name, rowid, name, path) VALUES('delete', old.id, old.name, old.path);
  INSERT INTO fts_name(rowid, name, path) VALUES (new.id, new.name, new.path);
END;
CREATE TRIGGER IF NOT EXISTS files_fts_ad AFTER DELETE ON files BEGIN
  INSERT INTO fts_name(fts_name, rowid, name, path) VALUES('delete', old.id, old.name, old.path);
END;
"""
FTS_OK = False   # create_app 时探测（老 SQLite 无 trigram 时自动退回 LIKE）


def get_db():
    conn = sqlite3.connect(DB_PATH, timeout=30)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA journal_mode=WAL")
    conn.execute("PRAGMA synchronous=NORMAL")
    return conn


# ---------------------------------------------------------------- 长路径安全
def _lp(p):
    """Windows 深层微信路径可能超 260 字符，读取时加 \\\\?\\ 前缀。"""
    p = os.path.abspath(p)
    if len(p) >= 240 and not p.startswith("\\\\?\\"):
        return "\\\\" if p.startswith("//") else "\\\\?\\" + p
    return p


def _clean_path(p):
    """用户粘贴的路径：去掉首尾空白与成对引号（右键"复制文件地址"自带引号）。"""
    p = (p or "").strip().strip('"').strip("'").strip()
    return os.path.normpath(p) if p else ""


def safe_stat(path):
    return os.stat(_lp(path))


def open_rb(path):
    return open(_lp(path), "rb")


# ---------------------------------------------------------------- 扫描状态
_scan_lock = threading.Lock()
SCAN = {
    "running": False, "phase": "idle", "root": "", "walked": 0, "found": 0,
    "added": 0, "updated": 0, "deleted": 0, "errors": 0,
    "started": None, "finished": None, "elapsed": 0,
    "extract_total": 0, "extract_done": 0,
}


def _scan_progress(**kw):
    with _scan_lock:
        SCAN.update(kw)


def scan_state():
    with _scan_lock:
        return dict(SCAN)


# ---------------------------------------------------------------- 正文提取
class Unsupported(Exception):
    pass


def extract_text(path, ext):
    """返回文本；不支持抛 Unsupported；其他异常向上抛。"""
    if ext == "pdf":
        from pypdf import PdfReader
        reader = PdfReader(_lp(path))
        parts = []
        for i, page in enumerate(reader.pages[:EXTRACT_PAGES]):
            try:
                parts.append(page.extract_text() or "")
            except Exception:
                continue
            if sum(len(p) for p in parts) > CONTENT_LIMIT:
                break
        return "\n".join(parts)[:CONTENT_LIMIT]

    if ext == "docx":
        import docx
        d = docx.Document(_lp(path))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for row in t.rows:
                parts.append(" ".join(c.text for c in row.cells))
        return "\n".join(parts)[:CONTENT_LIMIT]

    if ext in ("xlsx", "xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(_lp(path), read_only=True, data_only=True)
        parts = []
        for ws in wb.worksheets[:5]:
            parts.append(f"[{ws.title}]")
            for row in ws.iter_rows(max_row=500, values_only=True):
                line = " ".join(str(c) for c in row if c is not None)
                if line.strip():
                    parts.append(line)
                if sum(len(p) for p in parts) > CONTENT_LIMIT:
                    break
        wb.close()
        return "\n".join(parts)[:CONTENT_LIMIT]

    if ext in TEXT_EXTS:
        with open_rb(path) as f:
            raw = f.read(CONTENT_LIMIT)
        for enc in ("utf-8", "gbk", "utf-16"):
            try:
                return raw.decode(enc)
            except (UnicodeDecodeError, UnicodeError):
                continue
        return raw.decode("utf-8", errors="replace")

    raise Unsupported(ext)


# ---------------------------------------------------------------- 哈希
def md5_of(path):
    h = hashlib.md5()
    with open_rb(path) as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _shell_thumb(path, w, cache):
    """Windows 资源管理器同款缩略图（视频/HEIC 等交给系统编解码器）。
    ctypes 直调 Shell COM + GDI，无 pywin32/opencv 依赖；无缩略图时返回 False。"""
    try:
        import ctypes
        from ctypes import wintypes
        from PIL import Image

        class GUID(ctypes.Structure):
            _fields_ = [("d1", ctypes.c_ulong), ("d2", ctypes.c_ushort),
                        ("d3", ctypes.c_ushort), ("d4", ctypes.c_ubyte * 8)]
        class SIZE(ctypes.Structure):
            _fields_ = [("cx", ctypes.c_long), ("cy", ctypes.c_long)]
        class BITMAP(ctypes.Structure):
            _fields_ = [("bmType", ctypes.c_long), ("bmWidth", ctypes.c_long),
                        ("bmHeight", ctypes.c_long), ("bmWidthBytes", ctypes.c_long),
                        ("bmPlanes", ctypes.c_ushort), ("bmBitsPixel", ctypes.c_ushort),
                        ("bmBits", ctypes.c_void_p)]
        class BMIH(ctypes.Structure):
            _fields_ = [("biSize", ctypes.c_uint32), ("biWidth", ctypes.c_long),
                        ("biHeight", ctypes.c_long), ("biPlanes", ctypes.c_ushort),
                        ("biBitCount", ctypes.c_ushort), ("biCompression", ctypes.c_uint32),
                        ("biSizeImage", ctypes.c_uint32), ("biXPelsPerMeter", ctypes.c_long),
                        ("biYPelsPerMeter", ctypes.c_long), ("biClrUsed", ctypes.c_uint32),
                        ("biClrImportant", ctypes.c_uint32)]
        class BMI(ctypes.Structure):
            _fields_ = [("bmiHeader", BMIH), ("bmiColors", ctypes.c_uint32 * 3)]

        iid = GUID()
        ctypes.windll.ole32.CLSIDFromString(
            ctypes.c_wchar_p("{bcc18b79-ba16-442f-80c4-8a59c30c463b}"), ctypes.byref(iid))
        ole32 = ctypes.windll.ole32
        rc = ole32.CoInitializeEx(None, 0)          # COINIT_MULTITHREADED
        changed = (rc == 0x80010106)                # 线程已有其它套间模式 → 忽略
        if rc not in (0, 0x80010106):
            return False
        try:
            ppv = ctypes.c_void_p()
            rc = ctypes.windll.shell32.SHCreateItemFromParsingName(
                ctypes.c_wchar_p(path), None, ctypes.byref(iid), ctypes.byref(ppv))
            if rc != 0 or not ppv.value:
                return False
            vtbl = (ctypes.c_void_p * 8).from_address(
                ctypes.cast(ppv, ctypes.POINTER(ctypes.c_void_p)).contents.value)
            # IShellItemImageFactory::GetImage 位于虚表槽 3（vtbl[i] 本身就是函数地址）
            GetImage = ctypes.WINFUNCTYPE(
                ctypes.c_long, ctypes.c_void_p, SIZE, ctypes.c_uint,
                ctypes.POINTER(ctypes.c_void_p), ctypes.POINTER(SIZE))(vtbl[3])
            hbm = ctypes.c_void_p()
            size = SIZE()
            # SIIGBF_THUMBNAILONLY(0x8) | SIIGBF_BIGGERSIZEOK(0x1)
            if GetImage(ppv, SIZE(w, w), 0x8 | 0x1, ctypes.byref(hbm), ctypes.byref(size)) != 0 \
                    or not hbm.value:
                return False
            bm = BITMAP()
            ctypes.windll.gdi32.GetObjectW(hbm, ctypes.sizeof(BITMAP), ctypes.byref(bm))
            if bm.bmWidth <= 0 or bm.bmHeight <= 0:
                return False
            bmi = BMI()
            bmi.bmiHeader.biSize = ctypes.sizeof(BMIH)
            bmi.bmiHeader.biWidth = bm.bmWidth
            bmi.bmiHeader.biHeight = -abs(bm.bmHeight)     # top-down
            bmi.bmiHeader.biPlanes = 1
            bmi.bmiHeader.biBitCount = 32
            bmi.bmiHeader.biCompression = 0                # BI_RGB
            buf = ctypes.create_string_buffer(bm.bmWidth * abs(bm.bmHeight) * 4)
            hdc = ctypes.windll.user32.GetDC(0)
            try:
                lines = ctypes.windll.gdi32.GetDIBits(
                    hdc, hbm, 0, abs(bm.bmHeight), buf, ctypes.byref(bmi), 0)
            finally:
                ctypes.windll.user32.ReleaseDC(0, hdc)
                ctypes.windll.gdi32.DeleteObject(hbm)
            if lines <= 0:
                return False
            im = Image.frombuffer("RGBA", (bm.bmWidth, abs(bm.bmHeight)),
                                  buf.raw, "raw", "BGRA", 0, 1).convert("RGB")
            im.thumbnail((w, w))
            im.save(cache, "JPEG", quality=80)
            return True
        finally:
            if not changed:
                ole32.CoUninitialize()
    except Exception:
        return False


def _cv2_thumb(path, w, cache):
    """OpenCV 截帧兜底（本机装有 opencv-python 时可用）。"""
    try:
        import cv2
        cap = cv2.VideoCapture(_lp(path))
        if not cap.isOpened():
            # 个别 OpenCV/路径组合打不开时，退回 Windows 8.3 短路径
            try:
                import ctypes
                buf = ctypes.create_unicode_buffer(520)
                if ctypes.windll.kernel32.GetShortPathNameW(path, buf, 520):
                    cap = cv2.VideoCapture(buf.value)
            except Exception:
                pass
        fps = cap.get(cv2.CAP_PROP_FPS) or 25
        target = max(1, min(int(fps * 1.0), 120))
        got = False
        for _ in range(target):
            if cap.grab():
                got = True
            else:
                break
        ok = False
        if got:
            ok, frame = cap.retrieve()
            ok = ok and frame is not None
        if ok:
            # cv2.imwrite 不支持中文目录，改走 imencode + 文件句柄
            ok2, enc = cv2.imencode(".jpg", frame, [cv2.IMWRITE_JPEG_QUALITY, 80])
            if ok2:
                with open(cache, "wb") as fh:
                    fh.write(enc.tobytes())
                ok = True
            else:
                ok = False
        cap.release()
        return ok
    except Exception:
        return False


# ---------------------------------------------------------------- 扫描（增量）
def _record_error(conn, path, reason, phase):
    conn.execute("INSERT INTO errors(path, reason, phase, ts) VALUES(?,?,?,?)",
                 (path, str(reason)[:300], phase, time.time()))


def _deepen_pass(conn):
    """深化阶段：提取正文 + 哈希（扫描线程与启动续跑共用）。
    提速两把刀（借鉴 Everything 的思路——只做值得做的 I/O）：
    1. 同尺寸预筛：大小唯一的文件不可能是完全重复，跳过 MD5（通常省掉 90%+）；
    2. 线程池并行读盘：MD5/读文件是 I/O 等待，4 线程重叠等待时间。
    游标分批取，避免一次 fetchall 几十万行阻塞启动期响应。"""
    from concurrent.futures import ThreadPoolExecutor
    cur = conn.execute(
        "SELECT id, path, ext, size, content_status, hash_status FROM files "
        "WHERE content_status='pending' OR hash_status='pending'")
    total = conn.execute(
        "SELECT COUNT(*) c FROM files WHERE content_status='pending' OR hash_status='pending'").fetchone()["c"]
    dup_sizes = {r["size"] for r in conn.execute(
        "SELECT size FROM files WHERE size<=? GROUP BY size HAVING COUNT(*)>1", (HASH_LIMIT,))}
    _scan_progress(phase="extract", extract_total=total, extract_done=0)

    def process_one(r):
        """纯 I/O：读文件抽正文/算哈希，不碰数据库（线程安全）。"""
        out = {"id": r["id"], "content": None, "cstatus": None, "hash": None, "hstatus": None}
        if r["content_status"] == "pending":
            try:
                out["content"] = extract_text(r["path"], r["ext"])
                out["cstatus"] = "done"
            except Unsupported:
                out["cstatus"] = "none"
            except Exception:
                out["cstatus"] = "failed"
        if r["hash_status"] == "pending":
            if r["size"] > HASH_LIMIT:
                out["hstatus"] = "too_big"
            elif r["size"] not in dup_sizes:
                out["hstatus"] = "unique"           # 大小唯一 → 不可能是重复文件
            else:
                try:
                    out["hash"] = md5_of(r["path"])
                    out["hstatus"] = "done"
                except OSError:
                    out["hstatus"] = "failed"
        return out

    i = 0
    pool = ThreadPoolExecutor(max_workers=4)
    try:
        while True:
            rows = cur.fetchmany(2000)
            if not rows:
                break
            for r, out in zip(rows, pool.map(process_one, rows)):
                try:
                    if out["cstatus"]:
                        if out["cstatus"] == "failed":
                            _record_error(conn, r["path"], "正文提取失败", "extract")
                        conn.execute("UPDATE files SET content=?, content_status=? WHERE id=?",
                                     (out["content"], out["cstatus"], out["id"]))
                    if out["hstatus"]:
                        if out["hstatus"] == "failed":
                            _record_error(conn, r["path"], "哈希计算失败（读取被占用？）", "hash")
                        conn.execute("UPDATE files SET hash=?, hash_status=? WHERE id=?",
                                     (out["hash"], out["hstatus"], out["id"]))
                except Exception as e:  # 单文件失败绝不中断整体
                    _record_error(conn, r["path"], e, "deepen")
                i += 1
                if i % 100 == 0:
                    time.sleep(0.002)          # 让出 GIL：扫描期间网页必须保持响应
                if i % 500 == 0:
                    conn.commit()              # 批量提交：避免高频 fsync 开销
                    _scan_progress(extract_done=i)
    finally:
        pool.shutdown(wait=False)
        cur.close()
        conn.commit()


def _low_priority():
    """把当前线程降到低于普通优先级，扫描不抢前台。"""
    try:
        import ctypes
        ctypes.windll.kernel32.SetThreadPriority(
            ctypes.windll.kernel32.GetCurrentThread(), -1)   # THREAD_PRIORITY_BELOW_NORMAL
    except Exception:
        pass


def _scan_worker(root_ids, root_paths):
    conn = get_db()
    _low_priority()
    t0 = time.time()
    _scan_progress(running=True, phase="walk", walked=0, found=0, added=0, updated=0,
                   deleted=0, errors=0, started=t0, finished=None, elapsed=0,
                   extract_total=0, extract_done=0)
    seen = {}          # path -> (size, mtime)
    def _walk_err(e):
        try:
            with conn:
                _record_error(conn, getattr(e, "filename", "?") or "?", e, "walk")
        except Exception:
            pass
    try:
        for root_id, root in zip(root_ids, root_paths):
            _scan_progress(root=root)
            for dirpath, dirnames, filenames in os.walk(root, onerror=_walk_err):
                dirnames[:] = [d for d in dirnames if d not in ("$RECYCLE.BIN", "System Volume Information")]
                for fn in filenames:
                    _scan_progress(walked=SCAN["walked"] + 1)
                    if SCAN["walked"] % 500 == 0:
                        time.sleep(0.001)          # 让出 GIL，保证网页请求能插进来
                    full = os.path.join(dirpath, fn)
                    try:
                        st = safe_stat(full)
                    except OSError as e:
                        with conn:
                            _record_error(conn, full, e, "walk")
                        _scan_progress(errors=SCAN["errors"] + 1)
                        continue
                    seen[os.path.normcase(full)] = (st.st_size, st.st_mtime)
                    row = conn.execute("SELECT id, size, mtime FROM files WHERE path=?",
                                       (full,)).fetchone()
                    if row and row["size"] == st.st_size and abs(row["mtime"] - st.st_mtime) < 1e-6:
                        continue
                    ext = os.path.splitext(fn)[1].lstrip(".").lower()
                    folder = os.path.relpath(dirpath, root).replace("\\", "/")
                    if folder == ".":
                        folder = ""
                    ts = time.time()
                    hstat = "pending" if st.st_size <= HASH_LIMIT else "too_big"
                    if row:
                        conn.execute(
                            "UPDATE files SET root_id=?, name=?, ext=?, category=?, size=?, "
                            "ctime=?, mtime=?, folder=?, content=NULL, content_status='pending', "
                            "hash=NULL, hash_status=?, scan_ts=? WHERE id=?",
                            (root_id, fn, ext, category_of(ext), st.st_size,
                             st.st_ctime, st.st_mtime, folder, hstat, ts, row["id"]))
                        _scan_progress(updated=SCAN["updated"] + 1)
                    else:
                        conn.execute(
                            "INSERT INTO files(root_id, path, name, ext, category, size, ctime, "
                            "mtime, folder, content, content_status, hash, hash_status, scan_ts) "
                            "VALUES(?,?,?,?,?,?,?,?,?,NULL,'pending',NULL,?,?)",
                            (root_id, full, fn, ext, category_of(ext), st.st_size,
                             st.st_ctime, st.st_mtime, folder, hstat, ts))
                        _scan_progress(found=SCAN["found"] + 1)
                    if (SCAN["found"] + SCAN["updated"]) % 500 == 0:
                        conn.commit()
            # 该 root 下已消失的文件 → 删索引
            prefix = os.path.normcase(root)
            gone = [r["id"] for r in conn.execute(
                "SELECT id, path FROM files WHERE root_id=?", (root_id,))
                if os.path.normcase(r["path"]) not in seen]
            if gone:
                conn.executemany("DELETE FROM files WHERE id=?", [(i,) for i in gone])
                _scan_progress(deleted=SCAN["deleted"] + len(gone))
        conn.commit()

        # 深化：提取正文 + 哈希（复用函数）
        _deepen_pass(conn)
        _scan_progress(phase="done", running=False, elapsed=round(time.time() - t0, 1),
                       finished=time.time())
    except Exception as e:
        _record_error(conn, "SCAN", traceback.format_exc()[:500], "fatal")
        conn.commit()
        _scan_progress(phase="error", running=False, elapsed=round(time.time() - t0, 1),
                       finished=time.time())
    finally:
        conn.close()


def start_scan(root_ids, root_paths):
    if SCAN["running"]:
        return False
    _scan_progress(running=True, phase="starting", started=time.time(), finished=None)
    th = threading.Thread(target=_scan_worker, args=(root_ids, root_paths), daemon=True)
    th.start()
    return True


def start_deepen():
    """启动续跑：只补齐上次没提取完的正文/哈希，不重扫目录。"""
    if SCAN["running"]:
        return False

    def _worker():
        conn = get_db()
        _low_priority()
        t0 = time.time()
        _scan_progress(running=True, phase="extract", started=t0, finished=None,
                       walked=0, found=0, updated=0, deleted=0, errors=0,
                       extract_done=0, extract_total=0)
        try:
            _deepen_pass(conn)
            _scan_progress(phase="done", running=False,
                           elapsed=round(time.time() - t0, 1), finished=time.time())
        except Exception as e:
            _record_error(conn, "DEEPEN", traceback.format_exc()[:500], "fatal")
            conn.commit()
            _scan_progress(phase="error", running=False,
                           elapsed=round(time.time() - t0, 1), finished=time.time())
        finally:
            conn.close()

    th = threading.Thread(target=_worker, daemon=True)
    th.start()
    return True


# ---------------------------------------------------------------- 查询构建
SIZE_KEYS = ("min_size", "max_size")


def parse_filters(args):
    """从请求参数提取全部筛选条件 → dict。"""
    f = {
        "q": (args.get("q") or "").strip(),
        "scope": args.get("scope") or "all",                 # all/name/path/content
        "category": (args.get("category") or "").strip(),
        "root_id": args.get("root_id") or "",
        "folder": (args.get("folder") or "").strip(),
        "min_size": args.get("min_size") or "",
        "max_size": args.get("max_size") or "",
        "date_mode": "ctime" if args.get("date_mode") == "created" else "mtime",
        "date_type": args.get("date_type") or "any",          # any/year/month/day/range
        "year": args.get("year") or "",
        "month": args.get("month") or "",
        "day": args.get("day") or "",
        "date_from": args.get("date_from") or "",
        "date_to": args.get("date_to") or "",
        "sort": args.get("sort") or "mtime",
        "order": "asc" if args.get("order") == "asc" else "desc",
    }
    if f["scope"] not in ("all", "name", "path", "content"):
        f["scope"] = "all"
    if f["sort"] not in ("name", "size", "ctime", "mtime"):
        f["sort"] = "mtime"
    return f


def _mktime(y, m, d):
    return time.mktime((y, m, d, 0, 0, 0, 0, 0, -1))


def date_bounds(f):
    """由 year/month/day/range 计算 [start, end) epoch 区间。
    闰年、月末、跨年由 date+timedelta 自然处理。"""
    try:
        if f["date_type"] == "year" and f["year"]:
            y = int(f["year"])
            return _mktime(y, 1, 1), _mktime(y + 1, 1, 1)
        if f["date_type"] == "month" and f["year"] and f["month"]:
            y, m = int(f["year"]), int(f["month"])
            ny, nm = (y + 1, 1) if m == 12 else (y, m + 1)
            return _mktime(y, m, 1), _mktime(ny, nm, 1)
        if f["date_type"] == "day" and f["year"] and f["month"] and f["day"]:
            y, m, d = int(f["year"]), int(f["month"]), int(f["day"])
            nd = date(y, m, d) + timedelta(days=1)
            return _mktime(y, m, d), _mktime(nd.year, nd.month, nd.day)
        if f["date_type"] == "range" and f["date_from"]:
            y1, m1, d1 = (int(x) for x in f["date_from"].split("-"))
            start = _mktime(y1, m1, d1)
            if f["date_to"]:
                y2, m2, d2 = (int(x) for x in f["date_to"].split("-"))
                nd = date(y2, m2, d2) + timedelta(days=1)      # 含 date_to 当天（右开）
                end = _mktime(nd.year, nd.month, nd.day)
            else:
                end = time.time() + 86400
            return start, end
    except (ValueError, TypeError):
        pass
    return None


def _fts_phrase(w):
    return '"' + w.replace('"', '""') + '"'


def _word_conds(words, scope):
    """关键词列表 → (SQL 条件, 参数)。≥3 字符走 FTS trigram 索引，短词退回 LIKE。"""
    conds, params = [], []
    for w in words:
        alts = []
        if FTS_OK and len(w) >= 3:
            ph = _fts_phrase(w)
            if scope in ("all", "name"):
                alts.append("id IN (SELECT rowid FROM fts_name WHERE fts_name MATCH ?)")
                params.append(f"name:{ph}")
            if scope in ("all", "path"):
                alts.append("id IN (SELECT rowid FROM fts_name WHERE fts_name MATCH ?)")
                params.append(f"path:{ph}")
        else:
            like = f"%{_escape_like(w)}%"
            if scope in ("all", "name"):
                alts.append("name LIKE ? ESCAPE '\\'")
                params.append(like)
            if scope in ("all", "path"):
                alts.append("path LIKE ? ESCAPE '\\'")
                params.append(like)
        if scope in ("all", "content"):
            alts.append("(content LIKE ? ESCAPE '\\' AND content_status='done')")
            params.append(f"%{_escape_like(w)}%")
        if alts:
            conds.append("(" + " OR ".join(alts) + ")")
    return conds, params


def build_where(f, exclude_date=False):
    where, params = [], []
    words = [w for w in re.split(r"\s+", f["q"]) if w] if f["q"] else []
    if words:
        conds, wparams = _word_conds(words, f["scope"])
        if conds:
            where.append("(" + " AND ".join(conds) + ")")
        params.extend(wparams)
    if f["category"]:
        where.append("category=?")
        params.append(f["category"])
    if f["root_id"]:
        where.append("root_id=?")
        params.append(f["root_id"])
    if f["folder"]:
        where.append("folder LIKE ? ESCAPE '\\'")
        params.append(f"%{f['folder']}%")
    if f["min_size"]:
        where.append("size>=?")
        params.append(int(float(f["min_size"])))
    if f["max_size"]:
        where.append("size<=?")
        params.append(int(float(f["max_size"])))
    if not exclude_date:
        bounds = date_bounds(f)
        if bounds:
            where.append(f"{f['date_mode']}>=? AND {f['date_mode']}<?")
            params.extend(bounds)
    return " AND ".join(where) or "1=1", params


def _escape_like(s):
    return s.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")


def _kw_words(f):
    return [w for w in re.split(r"\s+", f["q"]) if w] if f["q"] else []


def matched_fields(row, words, scope):
    """该行命中了哪些来源：name / path / content。"""
    hits = []
    name = row["name"].casefold()
    path = row["path"].casefold()
    content = (row["content"] or "").casefold()
    for w in words:
        w = w.casefold()
        if scope in ("all", "name") and w in name:
            hits.append("name")
        if scope in ("all", "path") and w in path:
            hits.append("path")
        if scope in ("all", "content") and row["content_status"] == "done" and w in content:
            hits.append("content")
    return sorted(set(hits))


def make_snippet(row, words):
    content = row["content"] or ""
    low = content.casefold()
    for w in words:
        i = low.find(w.casefold())
        if i >= 0:
            s = max(0, i - 40)
            frag = content[s:i + SNIPPET_LEN].replace("\n", " ")
            return ("…" if s > 0 else "") + frag + "…"
    return (content[:SNIPPET_LEN].replace("\n", " ") + "…") if content else ""


# ---------------------------------------------------------------- Flask
def create_app(db_path=None, data_dir=None):
    global DB_PATH, DATA_DIR, FTS_OK
    if db_path:
        DB_PATH = db_path
    if data_dir:
        DATA_DIR = data_dir
    os.makedirs(os.path.join(DATA_DIR, "thumb_cache"), exist_ok=True)
    os.makedirs(os.path.join(DATA_DIR, "exports"), exist_ok=True)
    conn = sqlite3.connect(DB_PATH, timeout=30)
    conn.executescript(SCHEMA)
    try:
        conn.executescript(FTS_DDL)
        FTS_OK = True
    except sqlite3.OperationalError:
        FTS_OK = False                       # SQLite 无 FTS5/trigram → 搜索退回 LIKE
    if FTS_OK:
        built = conn.execute("SELECT value FROM meta WHERE key='fts_built'").fetchone()
        if not built:
            conn.execute("INSERT INTO fts_name(fts_name) VALUES('rebuild')")
            conn.execute("INSERT OR REPLACE INTO meta(key, value) VALUES('fts_built','1')")
    conn.commit()
    conn.close()
    app = Flask(__name__, static_folder=STATIC_DIR, static_url_path="/static")

    def db():
        conn = sqlite3.connect(DB_PATH, timeout=30)
        conn.row_factory = sqlite3.Row
        conn.execute("PRAGMA journal_mode=WAL")
        return conn

    @app.get("/")
    def index():
        return send_file(os.path.join(STATIC_DIR, "index.html"))

    @app.get("/favicon.ico")
    def favicon():
        p = os.path.join(STATIC_DIR, "favicon.ico")
        return send_file(p) if os.path.isfile(p) else ("", 204)

    # ---- 状态
    @app.get("/api/state")
    def api_state():
        conn = db()
        total = conn.execute("SELECT COUNT(*) c FROM files").fetchone()["c"]
        done = conn.execute("SELECT COUNT(*) c FROM files WHERE content_status='done'").fetchone()["c"]
        pending = conn.execute(
            "SELECT COUNT(*) c FROM files WHERE content_status='pending'").fetchone()["c"]
        roots = []
        for r in conn.execute(
                "SELECT r.id, r.path, r.added_at, COUNT(f.id) n FROM roots r "
                "LEFT JOIN files f ON f.root_id=r.id GROUP BY r.id ORDER BY r.id"):
            roots.append(dict(r))
        conn.close()
        return jsonify(ok=True, total=total, content_done=done, content_pending=pending,
                       roots=roots, scan=scan_state(),
                       app=APP_NAME, version=VERSION, author=AUTHOR, copyright=COPYRIGHT)

    @app.get("/api/candidates")
    def api_candidates():
        """探测常见微信文件目录：注册表自定义保存位置 → 文档目录 → 各盘常见命名。"""
        cands = []

        def add(p):
            if not p:
                return
            p = _clean_path(p)
            if p and os.path.isdir(p):
                key = os.path.normcase(p)
                if key not in {os.path.normcase(c) for c in cands}:
                    cands.append(os.path.normpath(p))

        docs = os.path.join(os.path.expanduser("~"), "Documents")
        # 1) 注册表：微信自定义的文件保存位置（新旧版本键名都试一遍）
        try:
            import winreg
            probes = (
                ("Software\\Tencent\\WeChat", ("FileSavePath", "InstallPath")),
                ("Software\\Tencent\\xwechat", ("FileSavePath", "InstallPath", "DataPath")),
                ("Software\\Tencent\\xwechat_main", ("FileSavePath", "InstallPath", "DataPath")),
            )
            for key_path, val_names in probes:
                try:
                    k = winreg.OpenKey(winreg.HKEY_CURRENT_USER, key_path)
                    for v in val_names:
                        try:
                            val = winreg.QueryValueEx(k, v)[0]
                            if isinstance(val, str) and len(val) > 3:
                                add(val)
                                # 保存位置本身往往太上层，把下面的微信子目录也列出来
                                try:
                                    for sub in sorted(os.listdir(val))[:10]:
                                        sp = os.path.join(val, sub)
                                        if os.path.isdir(sp) and (
                                                sub.startswith("wxid") or "wechat" in sub.lower()
                                                or "xwechat" in sub.lower()):
                                            add(sp)
                                except OSError:
                                    pass
                        except OSError:
                            pass
                except OSError:
                    pass
        except Exception:
            pass
        # 2) 文档目录（新旧两代微信的默认位置）
        for base in ("WeChat Files", "xwechat_files"):
            p = os.path.join(docs, base)
            add(p)
            try:
                for sub in sorted(os.listdir(p))[:10]:
                    add(os.path.join(p, sub))
            except OSError:
                pass
        # 3) 各盘根目录的常见命名
        for letter in string.ascii_uppercase:
            root = f"{letter}:\\"
            if not os.path.isdir(root):
                continue
            for name in ("WeChat Files", "xwechat_files", "微信文件", "WeChatFiles"):
                add(os.path.join(root, name))
        return jsonify(ok=True, candidates=cands[:30])

    # ---- 根目录管理
    @app.post("/api/roots")
    def api_add_roots():
        data = request.get_json(force=True)
        paths = data.get("paths", [])
        conn = db()
        added, skipped, results = [], [], []
        ids, real = [], []
        for raw in paths:
            p = _clean_path(raw)
            if not p:
                continue
            if not os.path.isdir(p):
                results.append({"input": raw, "path": p, "ok": False, "reason": "目录不存在，请核对路径"})
                skipped.append({"path": p, "reason": "目录不存在"})
                continue
            norm = os.path.normpath(p)
            exist = conn.execute("SELECT id FROM roots WHERE path=?", (norm,)).fetchone()
            if exist:
                results.append({"input": raw, "path": norm, "ok": True, "reason": "已在列表中"})
                skipped.append({"path": norm, "reason": "已在列表中"})
                ids.append(exist["id"]); real.append(norm)
                continue
            cur = conn.execute("INSERT INTO roots(path, added_at) VALUES(?,?)",
                               (norm, time.time()))
            results.append({"input": raw, "path": norm, "ok": True, "reason": "已添加"})
            ids.append(cur.lastrowid); real.append(norm)
            added.append(norm)
        conn.commit(); conn.close()
        if ids and not scan_state()["running"]:
            start_scan(ids, real)
        return jsonify(ok=bool(ids), added=added, skipped=skipped, results=results)

    @app.post("/api/roots/remove")
    def api_remove_root():
        if scan_state()["running"]:
            return jsonify(ok=False, error="正在扫描，请稍后再移除"), 409
        rid = request.get_json(force=True).get("id")
        conn = db()
        conn.execute("DELETE FROM files WHERE root_id=?", (rid,))
        conn.execute("DELETE FROM roots WHERE id=?", (rid,))
        conn.commit(); conn.close()
        return jsonify(ok=True)

    @app.post("/api/rescan")
    def api_rescan():
        conn = db()
        rows = conn.execute("SELECT id, path FROM roots").fetchall()
        conn.close()
        if not rows:
            return jsonify(ok=False, error="尚未添加任何文件夹"), 400
        if scan_state()["running"]:
            return jsonify(ok=False, error="扫描已在进行中"), 409
        start_scan([r["id"] for r in rows], [r["path"] for r in rows])
        return jsonify(ok=True)

    @app.get("/api/scan/status")
    def api_scan_status():
        return jsonify(ok=True, scan=scan_state())

    # ---- 搜索
    @app.get("/api/search")
    def api_search():
        f = parse_filters(request.args)
        where, params = build_where(f)
        sort = f["sort"]
        order = f["order"].upper()
        page = max(1, int(request.args.get("page", 1)))
        page_size = min(200, max(10, int(request.args.get("page_size", 50))))
        conn = db()
        total = conn.execute(f"SELECT COUNT(*) c FROM files WHERE {where}", params).fetchone()["c"]
        rows = conn.execute(
            f"SELECT * FROM files WHERE {where} ORDER BY {sort} {order}, id ASC LIMIT ? OFFSET ?",
            params + [page_size, (page - 1) * page_size]).fetchall()
        conn.close()
        words = _kw_words(f)
        items = []
        for r in rows:
            d = dict(r)
            d.pop("content", None)
            row_full = r
            d["matched"] = matched_fields(row_full, words, f["scope"])
            d["snippet"] = make_snippet(row_full, words) if "content" in d.get("matched", []) else ""
            items.append(d)
        pages = (total + page_size - 1) // page_size
        return jsonify(ok=True, total=total, page=page, page_size=page_size, pages=pages, items=items)

    @app.get("/api/folders")
    def api_folders():
        conn = db()
        q = "SELECT DISTINCT folder FROM files WHERE folder<>''"
        params = []
        if request.args.get("root_id"):
            q += " AND root_id=?"
            params.append(request.args["root_id"])
        q += " ORDER BY folder LIMIT 500"
        rows = conn.execute(q, params).fetchall()
        conn.close()
        return jsonify(ok=True, folders=[r["folder"] for r in rows])

    # ---- 图表（exclude: date,type,folder,root 任一维度不参与当前筛选，实现点图联动）
    @app.get("/api/charts")
    def api_charts():
        f = parse_filters(request.args)
        exclude = set((request.args.get("exclude") or "").split(",")) & {"date", "type", "folder", "root"}
        col = f["date_mode"]

        def build(excludes):
            w, p = [], []
            words = [x for x in re.split(r"\s+", f["q"]) if x] if f["q"] else []
            if words:
                conds, wparams = _word_conds(words, f["scope"])
                if conds:
                    w.append("(" + " AND ".join(conds) + ")")
                p.extend(wparams)
            if f["category"] and "type" not in excludes:
                w.append("category=?"); p.append(f["category"])
            if f["root_id"] and "root" not in excludes:
                w.append("root_id=?"); p.append(f["root_id"])
            if f["folder"] and "folder" not in excludes:
                w.append("folder LIKE ? ESCAPE '\\'"); p.append(f"%{f['folder']}%")
            if f["min_size"]:
                w.append("size>=?"); p.append(int(float(f["min_size"])))
            if f["max_size"]:
                w.append("size<=?"); p.append(int(float(f["max_size"])))
            if "date" not in excludes:
                b = date_bounds(f)
                if b:
                    w.append(f"{col}>=? AND {col}<?"); p.extend(b)
            return " AND ".join(w) or "1=1", p

        conn = db()
        out = {}
        if "date" in exclude:
            out["by_year"] = out["by_month"] = out["daily"] = []
        else:
            # 单趟整数分桶（按天）→ 月/年在 Python 聚合；避免 3 次 strftime 全表扫
            w, p = build({"date"})
            local = time.localtime()
            off = -(time.altzone if (time.daylight and local.tm_isdst) else time.timezone)
            rows = conn.execute(
                f"SELECT CAST(({col}+{int(off)})/86400 AS INTEGER) d, COUNT(*) c "
                f"FROM files WHERE {w} GROUP BY d", p).fetchall()
            daily, by_month, by_year = {}, {}, {}
            for r in rows:
                ymd = time.strftime("%Y-%m-%d", time.gmtime(r["d"] * 86400))
                daily[ymd] = r["c"]
                ym, yr = ymd[:7], ymd[:4]
                by_month[ym] = by_month.get(ym, 0) + r["c"]
                by_year[yr] = by_year.get(yr, 0) + r["c"]
            out["daily"] = [{"k": k, "c": c} for k, c in sorted(daily.items())]
            out["by_month"] = [{"k": k, "c": c} for k, c in sorted(by_month.items())]
            out["by_year"] = [{"k": k, "c": c} for k, c in sorted(by_year.items())]
        w, p = build({"type"})
        out["by_type"] = [dict(r) for r in conn.execute(
            f"SELECT category k, COUNT(*) c, SUM(size) s FROM files WHERE {w} "
            "GROUP BY category ORDER BY c DESC", p)]
        w, p = build({"folder"})
        raw = conn.execute(
            f"SELECT folder, COUNT(*) c, SUM(size) s FROM files WHERE {w} "
            "GROUP BY folder ORDER BY s DESC LIMIT 400", p).fetchall()
        merged = {}
        for r in raw:
            top = (r["folder"] or "(根目录)").split("/")[0]
            m = merged.setdefault(top, {"k": top, "c": 0, "s": 0})
            m["c"] += r["c"]; m["s"] += r["s"] or 0
        out["by_folder"] = sorted(merged.values(), key=lambda x: -x["s"])[:30]
        conn.close()
        return jsonify(ok=True, **out)

    # ---- 详情 / 预览 / 缩略图
    @app.get("/api/file/<int:fid>")
    def api_file(fid):
        conn = db()
        r = conn.execute("SELECT * FROM files WHERE id=?", (fid,)).fetchone()
        dup = None
        if r and r["hash"]:
            dup = conn.execute("SELECT COUNT(*) c FROM files WHERE hash=?", (r["hash"],)).fetchone()["c"]
        conn.close()
        if not r:
            abort(404)
        d = dict(r)
        d["dup_group_size"] = dup
        return jsonify(ok=True, file=d)

    @app.get("/api/preview/<int:fid>")
    def api_preview(fid):
        conn = db()
        r = conn.execute("SELECT path, ext FROM files WHERE id=?", (fid,)).fetchone()
        conn.close()
        if not r or not os.path.isfile(r["path"]):
            abort(404)
        mime = mimetypes.guess_type(r["path"])[0] or "application/octet-stream"
        return send_file(r["path"], mimetype=mime, conditional=True)

    IMG_EXTS = {"jpg", "jpeg", "png", "gif", "bmp", "webp", "tiff", "tif"}
    VID_EXTS = {"mp4", "avi", "mov", "mkv", "flv", "wmv", "ts", "3gp", "m4v", "webm"}

    @app.get("/api/thumb/<int:fid>")
    def api_thumb(fid):
        from PIL import Image
        conn = db()
        r = conn.execute("SELECT path, ext, mtime FROM files WHERE id=?", (fid,)).fetchone()
        conn.close()
        if not r:
            abort(404)
        w = min(640, max(80, int(request.args.get("w", 320))))
        cache = os.path.join(DATA_DIR, "thumb_cache", f"{fid}_{int(r['mtime'])}_{w}.jpg")
        if os.path.isfile(cache):
            return send_file(cache, mimetype="image/jpeg", max_age=3600)
        ok = False
        try:
            if r["ext"] in IMG_EXTS:
                im = Image.open(_lp(r["path"]))
                im = im.convert("RGB")
                im.thumbnail((w, w))
                im.save(cache, "JPEG", quality=80)
                ok = True
            elif r["ext"] in VID_EXTS or r["ext"] in ("heic", "heif"):
                # 首选 Windows 资源管理器同款缩略图（无 OpenCV 依赖，打包后体积小）
                ok = _shell_thumb(r["path"], w, cache)
                if not ok and r["ext"] in VID_EXTS:
                    ok = _cv2_thumb(r["path"], w, cache)
        except Exception:
            ok = False
        if ok and os.path.isfile(cache):
            return send_file(cache, mimetype="image/jpeg", max_age=3600)
        abort(404)

    # ---- 打开原文件
    @app.post("/api/open")
    def api_open():
        fid = request.get_json(force=True).get("id")
        conn = db()
        r = conn.execute("SELECT path FROM files WHERE id=?", (fid,)).fetchone()
        conn.close()
        if not r:
            return jsonify(ok=False, error="索引中不存在该文件"), 404
        if not os.path.isfile(r["path"]):
            return jsonify(ok=False, error="文件已不在磁盘上，请重扫索引"), 404
        try:
            os.startfile(r["path"])  # 只读打开，交给系统默认程序
            return jsonify(ok=True)
        except OSError as e:
            return jsonify(ok=False, error=str(e)), 500

    # ---- 重复 / 同名
    @app.get("/api/duplicates")
    def api_duplicates():
        conn = db()
        # 单遍扫描哈希索引（有序），凑满 200 组即停，避免逐组查询
        groups, prev, bucket = [], None, []
        for row in conn.execute(
                "SELECT id, name, path, size, mtime, ext, category, hash FROM files "
                "WHERE hash IS NOT NULL ORDER BY hash, path"):
            if row["hash"] != prev:
                if bucket:
                    if len(bucket) > 1:
                        groups.append({"hash": prev, "count": len(bucket),
                                       "total_size": sum(f["size"] for f in bucket),
                                       "files": bucket})
                        if len(groups) >= 200:
                            break
                    bucket = []
                prev = row["hash"]
            bucket.append(dict(row))
        if len(groups) < 200 and bucket and len(bucket) > 1:
            groups.append({"hash": prev, "count": len(bucket),
                           "total_size": sum(f["size"] for f in bucket), "files": bucket})
        conn.close()
        return jsonify(ok=True, groups=groups)

    @app.get("/api/same-names")
    def api_same_names():
        conn = db()
        groups = []
        for g in conn.execute(
                "SELECT name, COUNT(*) c FROM files GROUP BY name HAVING c>1 "
                "ORDER BY c DESC, name LIMIT 200"):
            files = [dict(x) for x in conn.execute(
                "SELECT id, name, path, size, mtime, ext, category FROM files WHERE name=? "
                "ORDER BY path", (g["name"],))]
            groups.append({"name": g["name"], "count": g["c"], "files": files})
        conn.close()
        return jsonify(ok=True, groups=groups)

    # ---- 导出
    @app.get("/api/export")
    def api_export():
        fmt = request.args.get("format", "csv")
        f = parse_filters(request.args)
        where, params = build_where(f)
        conn = db()
        rows = conn.execute(
            f"SELECT * FROM files WHERE {where} ORDER BY {f['sort']} {f['order'].upper()}, id "
            "LIMIT 50000", params).fetchall()
        conn.close()
        stamp = time.strftime("%Y%m%d_%H%M%S")
        os.makedirs(os.path.join(DATA_DIR, "exports"), exist_ok=True)
        if fmt == "md":
            out = os.path.join(DATA_DIR, "exports", f"微信文件地图_{stamp}.md")
            with open(out, "w", encoding="utf-8") as fh:
                fh.write(f"# 微信文件地图 · 导出结果\n\n- 导出时间：{time.strftime('%Y-%m-%d %H:%M:%S')}\n")
                fh.write(f"- 筛选条件：关键词「{f['q'] or '无'}」 类型「{f['category'] or '全部'}」"
                         f" 时间模式「{'创建' if f['date_mode']=='ctime' else '修改'}·{f['date_type']}」\n")
                fh.write(f"- 共 {len(rows)} 个文件\n\n")
                fh.write("| 文件名 | 类型 | 大小 | 创建时间 | 修改时间 | 完整路径 |\n|---|---|---|---|---|---|\n")
                for r in rows:
                    fh.write(f"| {r['name']} | {r['category']} | {r['size']/1024:.0f} KB | "
                             f"{time.strftime('%Y-%m-%d %H:%M', time.localtime(r['ctime']))} | "
                             f"{time.strftime('%Y-%m-%d %H:%M', time.localtime(r['mtime']))} | {r['path']} |\n")
        else:
            import csv
            out = os.path.join(DATA_DIR, "exports", f"微信文件地图_{stamp}.csv")
            with open(out, "w", encoding="utf-8-sig", newline="") as fh:   # utf-8-sig 防 Excel 中文乱码
                wcsv = csv.writer(fh)
                wcsv.writerow(["文件名", "类型", "扩展名", "大小(KB)", "创建时间", "修改时间",
                               "所在文件夹", "完整路径"])
                for r in rows:
                    wcsv.writerow([r["name"], r["category"], r["ext"], round(r["size"] / 1024, 1),
                                   time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(r["ctime"])),
                                   time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(r["mtime"])),
                                   r["folder"], r["path"]])
        return send_file(out, as_attachment=True)

    @app.get("/api/errors")
    def api_errors():
        conn = db()
        rows = [dict(r) for r in conn.execute(
            "SELECT * FROM errors ORDER BY id DESC LIMIT 100")]
        conn.close()
        return jsonify(ok=True, errors=rows)

    return app


app = create_app()

if __name__ == "__main__":
    os.makedirs(DATA_DIR, exist_ok=True)
    import socket
    port = None
    for p in range(PORT, PORT + 9):
        with socket.socket() as s:
            try:
                s.bind(("127.0.0.1", p))
                port = p
                break
            except OSError:
                continue
    if port is None:
        sys.exit(f"{PORT}-{PORT+8} 端口均被占用，无法启动")
    if FROZEN:
        # 双击 exe：起服务后自动打开浏览器
        import webbrowser
        threading.Timer(1.5, lambda: webbrowser.open(f"http://127.0.0.1:{port}")).start()
    # 上次没提取完的正文/哈希 → 后台续跑
    conn = sqlite3.connect(DB_PATH)
    pend = conn.execute(
        "SELECT COUNT(*) FROM files WHERE content_status='pending' "
        "OR (hash_status='pending' AND size<=?)", (HASH_LIMIT,)).fetchone()[0]
    conn.close()
    if pend:
        start_deepen()
        print(f"续跑：{pend} 个文件待提取正文/哈希")
    print(f"微信文件地图 → http://127.0.0.1:{port}")
    app.run(host="127.0.0.1", port=port, debug=False, threaded=True)
