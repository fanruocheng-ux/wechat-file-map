# -*- coding: utf-8 -*-
"""微信文件地图 测试套件
覆盖：扫描建索引、增量重扫、只读安全、日期边界（闰年/跨年/月末/同日多文件）、
组合筛选、内容搜索与命中来源标注、重复/同名识别、图表聚合、导出。
运行：python test_wxmap.py
"""
import csv
import io
import os
import sys
import time
import shutil
import hashlib
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import server

PY = os.path.dirname(os.path.abspath(__file__))


def _mtime(path, y, m, d, hh=0, mm=0):
    ts = time.mktime((y, m, d, hh, mm, 0, 0, 0, -1))
    os.utime(path, (ts, ts))


def _write(path, data):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    mode = "wb" if isinstance(data, bytes) else "w"
    with open(path, mode, encoding=None if isinstance(data, bytes) else "utf-8") as f:
        f.write(data)


def make_pdf(path, text):
    """用 reportlab 生成真实可提取文本的 PDF。"""
    from reportlab.pdfgen import canvas
    os.makedirs(os.path.dirname(path), exist_ok=True)
    c = canvas.Canvas(path)
    c.setFont("Helvetica", 12)
    c.drawString(72, 700, text)
    c.save()


class WxMapTestBase(unittest.TestCase):
    """公共夹具：一个仿微信目录结构的临时文件夹 + Flask test_client。"""

    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="wxmap_test_")
        self.db = os.path.join(self.tmp, "index.db")
        self.data = os.path.join(self.tmp, "data")
        self.root = os.path.join(self.tmp, "wxroot")
        self.build_fixture()
        self.snap_before = self.snapshot()
        server.create_app(db_path=self.db, data_dir=self.data)
        self.client = server.app.test_client()

    def tearDown(self):
        try:
            shutil.rmtree(self.tmp, ignore_errors=True)
        except Exception:
            pass

    # ---------- 夹具 ----------
    def build_fixture(self):
        r = self.root
        # 正文提取类
        make_pdf(os.path.join(r, "FileStorage", "File", "2024-02", "合同-扫描件.pdf"),
                 "contract amount 38000 yuan")
        import openpyxl
        wb = openpyxl.Workbook(); ws = wb.active
        ws["A1"] = "Item"; ws["B1"] = "Notebook"; ws["A2"] = "Qty"; ws["B2"] = 12
        p = os.path.join(r, "FileStorage", "File", "2024-03", "报价单.xlsx")
        os.makedirs(os.path.dirname(p), exist_ok=True)
        wb.save(p)
        import docx
        d = docx.Document()
        d.add_paragraph("antelope project kickoff meeting minutes, start next week")
        dp = os.path.join(r, "FileStorage", "File", "2024-03", "会议纪要.docx")
        os.makedirs(os.path.dirname(dp), exist_ok=True)
        d.save(dp)
        _write(os.path.join(r, "FileStorage", "File", "2024-03", "说明.md"),
               "# 说明书\n快速检索 leap year 支持\n")
        # 图片
        from PIL import Image
        ip = os.path.join(r, "FileStorage", "Image", "照片.png")
        os.makedirs(os.path.dirname(ip), exist_ok=True)
        Image.new("RGB", (120, 90), (200, 30, 30)).save(ip)
        # 损坏 PDF（提取失败 → errors 记录）
        _write(os.path.join(r, "FileStorage", "File", "2024-07", "损坏.pdf"), b"this is not a pdf at all")
        # 日期边界
        _write(os.path.join(r, "leap.txt"), "leap day file")
        _write(os.path.join(r, "nye.txt"), "new year eve")
        _write(os.path.join(r, "ny.txt"), "new year day")
        _write(os.path.join(r, "monthend.txt"), "month end file")
        _write(os.path.join(r, "same1.txt"), "x" * 45000)
        _write(os.path.join(r, "same2.txt"), "same day two")
        _write(os.path.join(r, "dup_a.txt"), "SAME-CONTENT-FOR-DUPLICATE-CHECK-1234567890")
        _write(os.path.join(r, "Sub", "dup_b.txt"), "SAME-CONTENT-FOR-DUPLICATE-CHECK-1234567890")
        _write(os.path.join(r, "同名文件.txt"), "AAA version")
        _write(os.path.join(r, "Other", "同名文件.txt"), "BBB version")
        _write(os.path.join(r, "big.txt"), "B" * 60000)
        # 时间（本地时间）
        _mtime(os.path.join(r, "FileStorage", "File", "2024-02", "合同-扫描件.pdf"), 2024, 5, 15, 10, 0)
        _mtime(os.path.join(r, "FileStorage", "File", "2024-03", "报价单.xlsx"), 2024, 5, 15, 10, 5)
        _mtime(os.path.join(r, "FileStorage", "File", "2024-03", "会议纪要.docx"), 2024, 5, 16, 11, 0)
        _mtime(os.path.join(r, "FileStorage", "File", "2024-03", "说明.md"), 2024, 5, 16, 11, 1)
        _mtime(os.path.join(r, "FileStorage", "Image", "照片.png"), 2024, 7, 1, 9, 0)
        _mtime(os.path.join(r, "FileStorage", "File", "2024-07", "损坏.pdf"), 2024, 7, 2, 9, 0)
        _mtime(os.path.join(r, "leap.txt"), 2024, 2, 29, 10, 0)      # 闰年 2-29
        _mtime(os.path.join(r, "nye.txt"), 2023, 12, 31, 23, 30)     # 跨年前夜
        _mtime(os.path.join(r, "ny.txt"), 2024, 1, 1, 0, 30)         # 跨年凌晨
        _mtime(os.path.join(r, "monthend.txt"), 2024, 4, 30, 23, 59) # 月末
        _mtime(os.path.join(r, "same1.txt"), 2024, 4, 30, 8, 0)      # 同日多文件
        _mtime(os.path.join(r, "same2.txt"), 2024, 4, 30, 12, 0)
        _mtime(os.path.join(r, "dup_a.txt"), 2024, 6, 1, 9, 0)
        _mtime(os.path.join(r, "Sub", "dup_b.txt"), 2024, 6, 1, 10, 0)
        _mtime(os.path.join(r, "同名文件.txt"), 2024, 6, 2, 9, 0)
        _mtime(os.path.join(r, "Other", "同名文件.txt"), 2024, 6, 2, 10, 0)
        _mtime(os.path.join(r, "big.txt"), 2024, 8, 1, 12, 0)

    def snapshot(self):
        """只读安全校验用的全量快照。"""
        out = {}
        for dp, _, fns in os.walk(self.root):
            for fn in fns:
                p = os.path.join(dp, fn)
                st = os.stat(p)
                with open(p, "rb") as f:
                    md5 = hashlib.md5(f.read()).hexdigest()
                out[os.path.normcase(p)] = (st.st_size, st.st_mtime_ns, md5)
        return out

    # ---------- 公共 ----------
    def scan(self, paths=None, rescan=False):
        if rescan:
            rv = self.client.post("/api/rescan")
        else:
            rv = self.client.post("/api/roots", json={"paths": paths or [self.root]})
        self.assertEqual(rv.status_code, 200, rv.get_data(as_text=True))
        deadline = time.time() + 120
        seen_running = False
        while time.time() < deadline:
            st = self.client.get("/api/scan/status").get_json()["scan"]
            if st["running"]:
                seen_running = True
            if seen_running and not st["running"] and st["phase"] in ("done", "error"):
                self.assertEqual(st["phase"], "done", st)
                return st
            time.sleep(0.1)
        self.fail("扫描超时")

    def search(self, **kw):
        return self.client.get("/api/search", query_string=kw).get_json()


class TestAddRoots(WxMapTestBase):
    def test_quoted_path_stripped(self):
        """右键"复制文件地址"自带引号 → 剥离后正常添加。"""
        j = self.client.post("/api/roots", json={"paths": [f'"{self.root}"']}).get_json()
        self.assertTrue(j["ok"])
        self.assertEqual(len(j["added"]), 1)

    def test_invalid_path_fails_loudly(self):
        """无效路径必须明确失败并给出原因，而不是静默成功。"""
        j = self.client.post("/api/roots", json={"paths": [r"F:\不存在的目录xyz"]}).get_json()
        self.assertFalse(j["ok"])
        self.assertIn("目录不存在", j["results"][0]["reason"])


class TestScanAndIndex(WxMapTestBase):
    def test_scan_counts_and_fields(self):
        st = self.scan()
        self.assertEqual(st["found"], 17)
        j = self.search(q="", page_size=100)
        self.assertEqual(j["total"], 17)
        one = self.search(q="合同-扫描件.pdf", scope="name")
        self.assertEqual(one["total"], 1)
        it = one["items"][0]
        for field in ("path", "name", "ext", "category", "size", "ctime", "mtime", "folder"):
            self.assertIn(field, it)
        self.assertEqual(it["ext"], "pdf")
        self.assertEqual(it["category"], "文档")
        self.assertIn("2024-02", it["folder"])

    def test_created_mode_smoke(self):
        self.scan()
        j = self.client.get("/api/search", query_string={"date_mode": "created", "page_size": 10}).get_json()
        self.assertTrue(j["ok"])
        self.assertEqual(j["total"], 17)


class TestIncrementalRescan(WxMapTestBase):
    def test_add_modify_delete(self):
        self.scan()
        # 改（尺寸变化）
        leap = os.path.join(self.root, "leap.txt")
        _write(leap, "leap day file with longer content now")
        _mtime(leap, 2024, 2, 29, 11, 0)
        # 删
        os.remove(os.path.join(self.root, "ny.txt"))
        # 增
        newf = os.path.join(self.root, "new.txt")
        _write(newf, "brand new file")
        _mtime(newf, 2024, 9, 1, 8, 0)
        st = self.scan(rescan=True)
        self.assertEqual(st["found"], 1)
        self.assertEqual(st["updated"], 1)
        self.assertEqual(st["deleted"], 1)
        j = self.search(q="ny.txt", scope="name")
        self.assertEqual(j["total"], 0)
        j = self.search(q="new.txt", scope="name")
        self.assertEqual(j["total"], 1)
        # 未变文件不重提取
        conn = server.sqlite3.connect(self.db)
        row = conn.execute("SELECT content_status FROM files WHERE name='说明.md'").fetchone()
        conn.close()
        self.assertEqual(row[0], "done")


class TestReadOnlySafety(WxMapTestBase):
    def test_no_source_file_touched(self):
        self.scan()
        # 触发缩略图与预览（读原文件）
        j = self.search(q="照片.png", scope="name")
        fid = j["items"][0]["id"]
        rv = self.client.get(f"/api/thumb/{fid}")
        self.assertEqual(rv.status_code, 200)
        self.assertEqual(rv.mimetype, "image/jpeg")
        rv = self.client.get(f"/api/preview/{fid}")
        self.assertEqual(rv.status_code, 200)
        self.assertEqual(self.snapshot(), self.snap_before,
                         "源文件被改动！只读承诺被破坏")


class TestDateBoundaries(WxMapTestBase):
    def setUp(self):
        super().setUp()
        self.scan()

    def q(self, **kw):
        kw.setdefault("date_mode", "modified")
        return self.search(**kw)["total"]

    def test_leap_feb29(self):
        self.assertEqual(self.q(date_type="day", year=2024, month=2, day=29), 1)
        self.assertEqual(self.q(date_type="month", year=2024, month=2), 1)
        self.assertEqual(self.q(date_type="day", year=2024, month=2, day=28), 0)

    def test_cross_year(self):
        self.assertEqual(self.q(date_type="range", date_from="2023-12-31", date_to="2024-01-01"), 2)
        self.assertEqual(self.q(date_type="range", date_from="2023-12-01", date_to="2024-03-01"), 3)
        self.assertEqual(self.q(date_type="year", year=2024), 16)
        self.assertEqual(self.q(date_type="year", year=2023), 1)

    def test_month_end_and_same_day(self):
        self.assertEqual(self.q(date_type="day", year=2024, month=4, day=30), 3)
        self.assertEqual(self.q(date_type="month", year=2024, month=4), 3)
        self.assertEqual(self.q(date_type="month", year=2024, month=5), 4)
        self.assertEqual(self.q(date_type="range", date_from="2024-05-15", date_to="2024-05-15"), 2)

    def test_charts_daily(self):
        j = self.client.get("/api/charts", query_string={"date_mode": "modified"}).get_json()
        daily = {d["k"]: d["c"] for d in j["daily"]}
        self.assertEqual(daily.get("2024-02-29"), 1)
        self.assertEqual(daily.get("2024-04-30"), 3)
        by_type = {t["k"]: t["c"] for t in j["by_type"]}
        self.assertEqual(by_type.get("表格"), 1)
        self.assertTrue(j["by_folder"])


class TestCombinedFilters(WxMapTestBase):
    def test_keyword_type_date(self):
        self.scan()
        self.assertEqual(self.search(category="文档", date_mode="modified",
                                     date_type="range", date_from="2024-05-15",
                                     date_to="2024-05-16")["total"], 3)

    def test_size_range(self):
        self.scan()
        self.assertEqual(self.search(min_size=40000, max_size=80000)["total"], 2)  # same1(45K)+big(60K)
        self.assertEqual(self.search(min_size=50000, max_size=80000)["total"], 1)
        self.assertEqual(self.search(min_size=40000, max_size=80000, date_mode="modified",
                                     date_type="month", year=2024, month=4)["total"], 1)

    def test_folder_and_root(self):
        self.scan()
        self.assertEqual(self.search(folder="Other")["total"], 1)
        self.assertEqual(self.search(folder="File/2024-03")["total"], 3)


class TestContentSearch(WxMapTestBase):
    def setUp(self):
        super().setUp()
        self.scan()

    def test_content_hit_labelled(self):
        j = self.search(q="antelope")
        self.assertEqual(j["total"], 1)
        self.assertIn("content", j["items"][0]["matched"])
        j = self.search(q="38000")
        self.assertEqual(j["total"], 1)
        self.assertEqual(j["items"][0]["ext"], "pdf")

    def test_scope_content_only(self):
        j = self.search(q="Notebook", scope="content")
        self.assertEqual(j["total"], 1)
        self.assertEqual(j["items"][0]["ext"], "xlsx")
        self.assertNotIn("name", j["items"][0]["matched"])

    def test_unextractable_still_findable_by_name(self):
        j = self.search(q="损坏.pdf", scope="name")
        self.assertEqual(j["total"], 1)
        j = self.search(q="pdf at all")   # 正文没被提取 → 内容搜不到
        self.assertEqual(j["total"], 0)
        # 失败有记录
        errs = self.client.get("/api/errors").get_json()["errors"]
        self.assertTrue(any(e["phase"] == "extract" for e in errs))


class TestDuplicates(WxMapTestBase):
    def setUp(self):
        super().setUp()
        self.scan()

    def test_hash_duplicates(self):
        j = self.client.get("/api/duplicates").get_json()
        self.assertEqual(len(j["groups"]), 1)
        g = j["groups"][0]
        self.assertEqual(g["count"], 2)
        names = {f["name"] for f in g["files"]}
        self.assertEqual(names, {"dup_a.txt", "dup_b.txt"})

    def test_same_names(self):
        j = self.client.get("/api/same-names").get_json()
        g = [x for x in j["groups"] if x["name"] == "同名文件.txt"]
        self.assertEqual(len(g), 1)
        self.assertEqual(g[0]["count"], 2)
        # 同名但内容不同 → 不进重复组
        dups = self.client.get("/api/duplicates").get_json()["groups"]
        all_dup_names = {f["name"] for grp in dups for f in grp["files"]}
        self.assertNotIn("同名文件.txt", all_dup_names)

    def test_unique_size_prefilter(self):
        """大小唯一的文件直接标记 unique（Everything 式预筛），不算 MD5。"""
        j = self.search(q="说明.md", scope="name")
        fid = j["items"][0]["id"]
        conn = server.sqlite3.connect(server.DB_PATH)
        row = conn.execute("SELECT hash_status FROM files WHERE id=?", (fid,)).fetchone()
        conn.close()
        self.assertEqual(row[0], "unique")
        j = self.search(q="dup_a.txt", scope="name")
        fid = j["items"][0]["id"]
        conn = server.sqlite3.connect(server.DB_PATH)
        row = conn.execute("SELECT hash, hash_status FROM files WHERE id=?", (fid,)).fetchone()
        conn.close()
        self.assertEqual(row[1], "done")
        self.assertTrue(row[0])


class TestExport(WxMapTestBase):
    def test_csv_and_md(self):
        self.scan()
        qs = {"date_mode": "modified", "date_type": "day", "year": 2024, "month": 2, "day": 29}
        rv = self.client.get("/api/export", query_string={**qs, "format": "csv"})
        self.assertEqual(rv.status_code, 200)
        rows = list(csv.reader(io.StringIO(rv.get_data(as_text=True))))
        self.assertEqual(len(rows), 2)  # 表头 + 1 行
        self.assertIn("leap.txt", rows[1][7])
        rv = self.client.get("/api/export", query_string={**qs, "format": "md"})
        text = rv.get_data(as_text=True)
        self.assertIn("共 1 个文件", text)
        self.assertIn("leap.txt", text)


if __name__ == "__main__":
    unittest.main(verbosity=2)
